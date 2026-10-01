"""Integration tests for Document Ingestion Pipeline and Status Lifecycle (Phase 7)."""

import io
from pathlib import Path
import pytest
import pypdf
import docx
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session
from sqlalchemy import select
from backend.app.main import app
from backend.app.db.session import get_db
from backend.app.models.document import Document, DocumentStatus
from backend.app.ingestion.service import IngestionService


@pytest.fixture
def client(db_session: Session) -> TestClient:
    """Provides a TestClient with database session overridden to advanced_rag_test."""
    def _override_get_db():
        try:
            yield db_session
        finally:
            pass

    app.dependency_overrides[get_db] = _override_get_db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


def create_user_and_get_token(client: TestClient, email: str, password: str = "SecurePass123!") -> str:
    """Helper to register and login a user, returning their JWT access token."""
    client.post("/api/v1/auth/register", json={
        "email": email,
        "password": password,
        "full_name": f"User {email.split('@')[0]}",
    })
    login_res = client.post("/api/v1/auth/login", json={
        "email": email,
        "password": password,
    })
    return login_res.json()["access_token"]


class TestDocumentIngestionLifecycle:
    """Test full upload -> background extraction -> status transition lifecycle."""

    def test_pdf_ingestion_lifecycle(self, client: TestClient, db_session: Session):
        token = create_user_and_get_token(client, "pdf_ingest_user@example.com")
        headers = {"Authorization": f"Bearer {token}"}

        # Create a valid multi-page PDF in memory
        writer = pypdf.PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.add_blank_page(width=100, height=100)
        writer.add_blank_page(width=100, height=100)
        pdf_bytes = io.BytesIO()
        writer.write(pdf_bytes)
        pdf_bytes.seek(0)

        # Upload document
        files = {"file": ("research_paper.pdf", pdf_bytes, "application/pdf")}
        response = client.post("/api/v1/documents/upload", headers=headers, files=files)
        assert response.status_code == 201
        doc_id = response.json()["id"]

        # Run ingestion processing
        user_res = client.get("/api/v1/users/me", headers=headers)
        user_id = user_res.json()["id"]

        ingestion_svc = IngestionService()
        processed_doc = ingestion_svc.process_document(
            document_id=doc_id,
            user_id=user_id,
            db=db_session,
        )

        assert processed_doc is not None
        assert processed_doc.status == DocumentStatus.COMPLETED
        assert processed_doc.page_count == 3
        assert processed_doc.error_message is None

        # Verify via GET /documents/{id} endpoint
        get_res = client.get(f"/api/v1/documents/{doc_id}", headers=headers)
        assert get_res.status_code == 200
        get_data = get_res.json()
        assert get_data["status"] == "COMPLETED"
        assert get_data["page_count"] == 3

    def test_docx_ingestion_lifecycle(self, client: TestClient, db_session: Session):
        token = create_user_and_get_token(client, "docx_ingest_user@example.com")
        headers = {"Authorization": f"Bearer {token}"}

        # Create valid docx in memory
        doc_obj = docx.Document()
        doc_obj.add_heading("Product Requirements", level=1)
        doc_obj.add_paragraph("High performance RAG architecture specifications.")
        docx_bytes = io.BytesIO()
        doc_obj.save(docx_bytes)
        docx_bytes.seek(0)

        files = {"file": ("requirements.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        response = client.post("/api/v1/documents/upload", headers=headers, files=files)
        assert response.status_code == 201
        doc_id = response.json()["id"]

        user_id = client.get("/api/v1/users/me", headers=headers).json()["id"]

        ingestion_svc = IngestionService()
        processed_doc = ingestion_svc.process_document(
            document_id=doc_id,
            user_id=user_id,
            db=db_session,
        )

        assert processed_doc is not None
        assert processed_doc.status == DocumentStatus.COMPLETED
        assert processed_doc.error_message is None

    def test_txt_ingestion_lifecycle(self, client: TestClient, db_session: Session):
        token = create_user_and_get_token(client, "txt_ingest_user@example.com")
        headers = {"Authorization": f"Bearer {token}"}

        txt_bytes = io.BytesIO("Chapter 1: Foundations of Artificial Intelligence\n\nIntroduction to modern neural models.".encode("utf-8"))
        files = {"file": ("book_chapter.txt", txt_bytes, "text/plain")}
        response = client.post("/api/v1/documents/upload", headers=headers, files=files)
        assert response.status_code == 201
        doc_id = response.json()["id"]

        user_id = client.get("/api/v1/users/me", headers=headers).json()["id"]

        ingestion_svc = IngestionService()
        processed_doc = ingestion_svc.process_document(
            document_id=doc_id,
            user_id=user_id,
            db=db_session,
        )

        assert processed_doc is not None
        assert processed_doc.status == DocumentStatus.COMPLETED

    def test_corrupt_file_ingestion_fails_safely(self, client: TestClient, db_session: Session):
        token = create_user_and_get_token(client, "corrupt_ingest_user@example.com")
        headers = {"Authorization": f"Bearer {token}"}

        # Upload a file with valid PDF header to pass Phase 6 upload check, but corrupted body
        corrupt_pdf_bytes = io.BytesIO(b"%PDF-1.7 Corrupt Stream without Trailer or xref")
        files = {"file": ("damaged.pdf", corrupt_pdf_bytes, "application/pdf")}
        response = client.post("/api/v1/documents/upload", headers=headers, files=files)
        assert response.status_code == 201
        doc_id = response.json()["id"]

        user_id = client.get("/api/v1/users/me", headers=headers).json()["id"]

        # Run ingestion
        ingestion_svc = IngestionService()
        processed_doc = ingestion_svc.process_document(
            document_id=doc_id,
            user_id=user_id,
            db=db_session,
        )

        assert processed_doc is not None
        assert processed_doc.status == DocumentStatus.FAILED
        assert processed_doc.error_message is not None
        # Security assertion: No absolute server path or secrets exposed in error_message
        assert "C:" not in processed_doc.error_message
        assert "Users" not in processed_doc.error_message

        # Verify API response
        get_res = client.get(f"/api/v1/documents/{doc_id}", headers=headers)
        assert get_res.status_code == 200
        assert get_res.json()["status"] == "FAILED"
        assert get_res.json()["error_message"] is not None

    def test_ingestion_isolation_cannot_process_other_user_document(self, client: TestClient, db_session: Session):
        # User A uploads doc
        token_a = create_user_and_get_token(client, "owner_a@example.com")
        headers_a = {"Authorization": f"Bearer {token_a}"}
        res_a = client.post(
            "/api/v1/documents/upload",
            headers=headers_a,
            files={"file": ("doc_a.txt", io.BytesIO(b"Confidential user A text"), "text/plain")},
        )
        doc_a_id = res_a.json()["id"]

        # User B tries to invoke ingestion on User A's document ID
        token_b = create_user_and_get_token(client, "intruder_b@example.com")
        user_b_id = client.get("/api/v1/users/me", headers={"Authorization": f"Bearer {token_b}"}).json()["id"]

        ingestion_svc = IngestionService()
        result = ingestion_svc.process_document(
            document_id=doc_a_id,
            user_id=user_b_id,
            db=db_session,
        )

        assert result is None  # Fails safely because user_b does not own doc_a

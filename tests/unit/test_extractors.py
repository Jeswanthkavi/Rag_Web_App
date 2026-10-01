"""Unit tests for PDF, DOCX, TXT extractors and TextNormalizer (Phase 7)."""

from pathlib import Path
import pytest
import pypdf
import docx
from backend.app.ingestion.pdf_extractor import PDFExtractor
from backend.app.ingestion.docx_extractor import DOCXExtractor
from backend.app.ingestion.txt_extractor import TXTExtractor
from backend.app.ingestion.normalizer import TextNormalizer
from backend.app.ingestion.service import IngestionService


@pytest.fixture
def temp_dir(tmp_path: Path) -> Path:
    """Provides a temporary directory for test files."""
    return tmp_path


class TestTextNormalizer:
    """Verify deterministic text normalization."""

    def test_unicode_normalization(self):
        # Full-width characters and ligatures
        raw = "Ｆｕｌｌｗｉｄｔｈ ａｎｄ ﬁ (ligature)"
        normalized = TextNormalizer.normalize(raw)
        assert normalized == "Fullwidth and fi (ligature)"

    def test_line_ending_standardization(self):
        raw = "Line 1\r\nLine 2\rLine 3\nLine 4"
        normalized = TextNormalizer.normalize(raw)
        assert normalized == "Line 1\nLine 2\nLine 3\nLine 4"

    def test_control_character_stripping(self):
        raw = "Valid text\x00\x08with\x1fcontrol\x7fchars\tpreserved tab"
        normalized = TextNormalizer.normalize(raw)
        # Note: \t is normalized to a clean single space separating words
        assert normalized == "Valid textwithcontrolchars preserved tab"

    def test_whitespace_and_paragraph_preservation(self):
        raw = "   Paragraph 1 with   multiple    spaces.   \n\n\n\n   Paragraph 2 after excessive newlines.   "
        normalized = TextNormalizer.normalize(raw)
        assert normalized == "Paragraph 1 with multiple spaces.\n\nParagraph 2 after excessive newlines."


class TestPDFExtractor:
    """Verify PDF extraction and page awareness."""

    def test_extract_valid_multipage_pdf(self, temp_dir: Path):
        # Create 2-page PDF in memory using pypdf
        writer = pypdf.PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.add_blank_page(width=200, height=200)
        
        pdf_path = temp_dir / "multipage.pdf"
        with open(pdf_path, "wb") as f:
            writer.write(f)

        extractor = PDFExtractor()
        doc = extractor.extract(pdf_path, "multipage.pdf")

        assert doc.filename == "multipage.pdf"
        assert doc.file_type == ".pdf"
        assert doc.total_pages == 2
        assert len(doc.pages) == 2
        assert doc.pages[0].page_number == 1
        assert doc.pages[1].page_number == 2

    def test_corrupt_pdf_raises_value_error(self, temp_dir: Path):
        corrupt_path = temp_dir / "corrupt.pdf"
        corrupt_path.write_bytes(b"%PDF-1.7 Corrupt Garbage Data without EOF")

        extractor = PDFExtractor()
        with pytest.raises(ValueError) as exc:
            extractor.extract(corrupt_path, "corrupt.pdf")
        assert "PDF" in str(exc.value)


class TestDOCXExtractor:
    """Verify DOCX extraction of paragraphs and tables."""

    def test_extract_valid_docx(self, temp_dir: Path):
        # Create a DOCX with paragraphs and a table
        doc = docx.Document()
        doc.add_heading("Architecture Overview", level=1)
        doc.add_paragraph("This is the first paragraph with technical details.")
        doc.add_paragraph("This is the second paragraph describing components.")
        
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Service"
        table.cell(0, 1).text = "Port"
        table.cell(1, 0).text = "FastAPI"
        table.cell(1, 1).text = "8000"

        docx_path = temp_dir / "sample.docx"
        doc.save(str(docx_path))

        extractor = DOCXExtractor()
        extracted = extractor.extract(docx_path, "sample.docx")

        assert extracted.filename == "sample.docx"
        assert extracted.file_type == ".docx"
        assert len(extracted.pages) == 1
        assert extracted.pages[0].page_number is None
        assert "Architecture Overview" in extracted.pages[0].text
        assert "technical details" in extracted.pages[0].text
        assert "FastAPI | 8000" in extracted.pages[0].text

    def test_corrupt_docx_raises_value_error(self, temp_dir: Path):
        corrupt_path = temp_dir / "corrupt.docx"
        corrupt_path.write_bytes(b"PK\x03\x04Corrupt zip content")

        extractor = DOCXExtractor()
        with pytest.raises(ValueError) as exc:
            extractor.extract(corrupt_path, "corrupt.docx")
        assert "DOCX" in str(exc.value) or "OpenXML" in str(exc.value)


class TestTXTExtractor:
    """Verify plain text file extraction."""

    def test_extract_utf8_txt(self, temp_dir: Path):
        txt_path = temp_dir / "document.txt"
        content = "Line 1 of document.\nLine 2 with Unicode: ü ö ä."
        txt_path.write_text(content, encoding="utf-8")

        extractor = TXTExtractor()
        extracted = extractor.extract(txt_path, "document.txt")

        assert extracted.filename == "document.txt"
        assert extracted.file_type == ".txt"
        assert len(extracted.pages) == 1
        assert extracted.pages[0].page_number is None
        assert "Line 1 of document." in extracted.pages[0].text
        assert "ü ö ä" in extracted.pages[0].text


class TestIngestionServiceIntegration:
    """Verify end-to-end extraction and normalization service."""

    def test_service_extract_and_normalize(self, temp_dir: Path):
        txt_path = temp_dir / "dirty.txt"
        txt_path.write_text("  Raw text with   spaces.  \r\n\r\n\r\nSecond line.  ", encoding="utf-8")

        service = IngestionService()
        doc = service.extract_and_normalize(txt_path, "dirty.txt", ".txt")

        assert doc.pages[0].text == "Raw text with spaces.\n\nSecond line."
